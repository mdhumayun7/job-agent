import json
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from config import ENRICHED_OUTPUT_FILE, JOBS_FILE, YOUR_SKILLS, PERSONAL_INFO, PROFILE_SUMMARY

def load_jobs():
    src = Path(ENRICHED_OUTPUT_FILE) if Path(ENRICHED_OUTPUT_FILE).exists() else Path(JOBS_FILE)
    with open(src,"r",encoding="utf-8") as f: return json.load(f)

def tailor_job(job):
    title   = job.get("title","")
    company = job.get("company","")
    matched = job.get("matched_skills",[]) or []
    missing = job.get("missing_skills",[]) or []
    top     = matched[:5] or YOUR_SKILLS[:5]
    name    = PERSONAL_INFO.get("name","Candidate")
    objective = (
        f"Motivated {PERSONAL_INFO.get('degree','') or 'engineering'} graduate seeking {title} role at {company}, "
        f"leveraging expertise in {', '.join(top[:3])} to deliver impactful solutions."
    )
    cover = (
        f"Dear Hiring Manager,\n\n"
        f"I am writing to express my strong interest in the {title} position at {company}. "
        f"As a {PROFILE_SUMMARY.split('.')[0].lower()}, I bring hands-on experience in "
        f"{', '.join(top[:3])}, making me a strong fit for this role.\n\n"
        f"I am confident in my ability to contribute from day one and am eager to grow "
        f"with your team. I look forward to discussing how my background aligns with your needs.\n\n"
        f"Warm regards,\n{name}"
    )
    ats_keywords = list(set(top + [s for s in YOUR_SKILLS if s.lower() in job.get('title','').lower()]))[:10]
    one_liner = f"Passionate {title} candidate with {len(matched)} matching skills including {', '.join(top[:2])}." if top else f"Eager fresher applying for {title} at {company}."
    return {
        "objective":      objective,
        "cover_letter":   cover,
        "ats_keywords":   ", ".join(ats_keywords),
        "one_liner":      one_liner,
        "missing_skills": ", ".join(missing[:5]),
    }

def save_excel(jobs, path="output/tailored_jobs.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tailored Applications"
    headers = ["Score","Title","Company","Location","Source","One-Liner","Objective","Cover Letter","ATS Keywords","Missing Skills","Apply Link"]
    hfill = PatternFill("solid", fgColor="4F46E5")
    hfont = Font(color="FFFFFF", bold=True)
    for col,h in enumerate(headers,1):
        c = ws.cell(row=1,column=col,value=h)
        c.fill = hfill; c.font = hfont
        c.alignment = Alignment(horizontal="center")
    for row,job in enumerate(jobs,2):
        sc = job.get("ai_match_score",job.get("match_score",0))
        color = "86EFAC" if sc>=60 else ("FDE68A" if sc>=30 else "FCA5A5")
        ws.cell(row=row,column=1,value=sc).fill = PatternFill("solid",fgColor=color)
        ws.cell(row=row,column=2,value=job.get("title",""))
        ws.cell(row=row,column=3,value=job.get("company",""))
        ws.cell(row=row,column=4,value=job.get("location",""))
        ws.cell(row=row,column=5,value=job.get("source_group",job.get("source","")))
        ws.cell(row=row,column=6,value=job.get("one_liner",""))
        ws.cell(row=row,column=7,value=job.get("objective","")).alignment = Alignment(wrap_text=True)
        ws.cell(row=row,column=8,value=job.get("cover_letter","")).alignment = Alignment(wrap_text=True)
        ws.cell(row=row,column=9,value=job.get("ats_keywords",""))
        ws.cell(row=row,column=10,value=job.get("missing_skills",""))
        ws.cell(row=row,column=11,value=job.get("apply_url",""))
    for col in ws.columns:
        mx = max((len(str(c.value or "")) for c in col),default=10)
        ws.column_dimensions[col[0].column_letter].width = min(mx+4,60)
    Path(path).parent.mkdir(exist_ok=True)
    wb.save(path)
    print(f"  Saved: {path}")

def main():
    jobs = load_jobs()
    top  = sorted(jobs, key=lambda j: j.get("ai_match_score",j.get("match_score",0)), reverse=True)[:50]
    print(f"\nTailoring resumes for top {len(top)} jobs...")
    tailored = []
    for i,job in enumerate(top,1):
        t = tailor_job(job)
        tailored.append({**job, **t})
        if i%10==0: print(f"  {i}/{len(top)} done")
    save_excel(tailored)
    print(f"Done! {len(tailored)} jobs tailored.")

if __name__ == "__main__":
    main()
