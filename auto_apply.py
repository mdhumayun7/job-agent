import json
import re
import sys
import time
sys.path.insert(0, r"C:\Users\HP\Desktop\job_agent")

from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font

RESUME_PATH = r"C:\Users\HP\Desktop\job_agent\resume\MD_Humayun_Resume.pdf"
JOBS_FILE = r"C:\Users\HP\Desktop\job_agent\output\jobs_found.json"
LOG_FILE = r"C:\Users\HP\Desktop\job_agent\output\apply_log.xlsx"

MY_INFO = {"name":"MD Humayun","email":"humayunrahi739@gmail.com","phone":"8789350894","location":"Surat, Gujarat","college":"SVNIT Surat","degree":"M.Tech Computer Science","cgpa":"8.90","experience":"0","linkedin":"linkedin.com/in/mdhumayun","github":"github.com/mdhumayun"}

applied_log = []

def show_approval_screen(jobs):
    print("\n" + "="*60)
    print("  JOBS WAITING FOR YOUR APPROVAL")
    print("="*60)
    approved = []
    for i, job in enumerate(jobs, 1):
        print(f"\n[{i}/{len(jobs)}]")
        print(f"  Title   : {job.get('title','')}")
        print(f"  Company : {job.get('company','')}")
        print(f"  Location: {job.get('location','')}")
        print(f"  Salary  : {job.get('salary','Not mentioned')}")
        print(f"  Source  : {job.get('source','')}")
        print(f"  Link    : {job.get('apply_url','')[:60]}")
        choice = input("  Apply karna hai? (y/n/q to quit): ").strip().lower()
        if choice == "q": break
        if choice == "y": approved.append(job)
    print(f"\nApproved: {len(approved)} jobs")
    return approved

def apply_linkedin(page, job):
    try:
        import json as _json, os as _os
        cookie_file = r"C:\Users\HP\Desktop\job_agent\browser_data\linkedin_cookies.json"
        if _os.path.exists(cookie_file):
            with open(cookie_file, "r") as cf:
                cookies = _json.load(cf)
            try:
                page.context.add_cookies(cookies)
            except: pass
        url = job.get("apply_url", "")
        if "linkedin.com" not in url: return False, "Not LinkedIn"
        page.goto(url, timeout=30000)
        page.wait_for_timeout(3000)
        easy_apply = page.query_selector(".jobs-apply-button--top-card button")
        if not easy_apply: easy_apply = page.query_selector("button.jobs-apply-button")
        if not easy_apply: return False, "No Easy Apply button"
        easy_apply.click()
        page.wait_for_timeout(2000)
        for _ in range(10):
            phone = page.query_selector("input[id*=phoneNumber]")
            if phone: phone.fill(MY_INFO["phone"])
            city = page.query_selector("input[id*=city]")
            if city: city.fill(MY_INFO["location"])
            resume_upload = page.query_selector("input[type=file]")
            if resume_upload: resume_upload.set_input_files(RESUME_PATH)
            next_btn = page.query_selector("button[aria-label='Continue to next step']")
            submit_btn = page.query_selector("button[aria-label='Submit application']")
            review_btn = page.query_selector("button[aria-label='Review your application']")
            if submit_btn:
                submit_btn.click()
                page.wait_for_timeout(2000)
                return True, "Applied successfully"
            elif review_btn: review_btn.click()
            elif next_btn: next_btn.click()
            else: break
            page.wait_for_timeout(1500)
        return False, "Could not complete form"
    except Exception as e:
        return False, str(e)

def apply_naukri(page, job):
    try:
        url = job.get("apply_url", "")
        if "naukri.com" not in url: return False, "Not Naukri"
        page.goto(url, timeout=30000)
        page.wait_for_timeout(3000)
        apply_btn = page.query_selector("button#apply-button")
        if not apply_btn: apply_btn = page.query_selector(".apply-button")
        if not apply_btn: return False, "No apply button"
        apply_btn.click()
        page.wait_for_timeout(2000)
        resume_upload = page.query_selector("input[type=file]")
        if resume_upload: resume_upload.set_input_files(RESUME_PATH)
        submit = page.query_selector("button[type=submit]")
        if submit:
            submit.click()
            page.wait_for_timeout(2000)
            return True, "Applied successfully"
        return False, "Submit button not found"
    except Exception as e:
        return False, str(e)

def apply_indeed(page, job):
    try:
        url = job.get("apply_url", "")
        if "indeed.com" not in url: return False, "Not Indeed"
        page.goto(url, timeout=30000)
        page.wait_for_timeout(3000)
        apply_btn = page.query_selector("button[id*=apply]")
        if not apply_btn: apply_btn = page.query_selector(".ia-IndeedApplyButton")
        if not apply_btn: return False, "No apply button"
        apply_btn.click()
        page.wait_for_timeout(2000)
        name = page.query_selector("input[name*=name]")
        if name: name.fill(MY_INFO["name"])
        email = page.query_selector("input[type=email]")
        if email: email.fill(MY_INFO["email"])
        phone = page.query_selector("input[type=tel]")
        if phone: phone.fill(MY_INFO["phone"])
        resume_upload = page.query_selector("input[type=file]")
        if resume_upload: resume_upload.set_input_files(RESUME_PATH)
        submit = page.query_selector("button[type=submit]")
        if submit:
            submit.click()
            page.wait_for_timeout(2000)
            return True, "Applied successfully"
        return False, "Submit not found"
    except Exception as e:
        return False, str(e)

def save_apply_log():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Apply Log"
    headers = ["Date","Title","Company","Source","Status","Message","Apply Link"]
    hf = PatternFill("solid", fgColor="1E293B")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = hf
        cell.font = Font(color="FFFFFF", bold=True)
    for row, log in enumerate(applied_log, 2):
        ws.cell(row=row, column=1, value=log["date"])
        ws.cell(row=row, column=2, value=log["title"])
        ws.cell(row=row, column=3, value=log["company"])
        ws.cell(row=row, column=4, value=log["source"])
        sc = ws.cell(row=row, column=5, value=log["status"])
        if log["status"] == "Applied": sc.fill = PatternFill("solid", fgColor="86EFAC")
        else: sc.fill = PatternFill("solid", fgColor="FCA5A5")
        ws.cell(row=row, column=6, value=log["message"])
        ws.cell(row=row, column=7, value=log["link"])
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len+4, 50)
    wb.save(LOG_FILE)
    print(f"Apply log saved: {LOG_FILE}")

def run_auto_apply(top_n=20, headless=False):
    with open(JOBS_FILE, "r", encoding="utf-8") as f:
        jobs = json.load(f)
    top_jobs = sorted(jobs, key=lambda x: x.get("match_score", 0), reverse=True)[:top_n]
    approved_jobs = show_approval_screen(top_jobs)
    if not approved_jobs:
        print("Koi job approve nahi ki!")
        return
    print(f"\nApplying to {len(approved_jobs)} jobs...")
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=headless)
        context = browser.new_context()
        page = context.new_page()
        for i, job in enumerate(approved_jobs, 1):
            source = job.get("source", "").lower()
            title = job.get("title", "")
            company = job.get("company", "")
            print(f"\n[{i}/{len(approved_jobs)}] Applying: {title} @ {company}")
            if "linkedin" in source: success, msg = apply_linkedin(page, job)
            elif "naukri" in source: success, msg = apply_naukri(page, job)
            elif "indeed" in source: success, msg = apply_indeed(page, job)
            else: success, msg = False, "Manual apply needed"
            status = "Applied" if success else "Failed"
            print(f"  Status: {status} — {msg}")
            applied_log.append({"date":datetime.now().strftime("%d-%m-%Y %H:%M"),"title":title,"company":company,"source":source,"status":status,"message":msg,"link":job.get("apply_url","")})
            time.sleep(2)
        browser.close()
    save_apply_log()
    applied = sum(1 for l in applied_log if l["status"] == "Applied")
    failed = len(applied_log) - applied
    print(f"\nDone! Applied: {applied} | Failed: {failed}")
    print(f"Log: {LOG_FILE}")

if __name__ == "__main__":
    run_auto_apply(top_n=20, headless=False)
