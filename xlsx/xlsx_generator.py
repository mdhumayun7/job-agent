"""
Multi-sheet professional XLSX generator.
Sheets: All Jobs, New Jobs, Fresher Jobs, Internships, CSE Jobs,
Closing Soon, Statistics.

Consumes a list of enriched Job dicts (output of parsers.enrich_job).
Never invents data -- every cell either comes from the job dict or is
"Not specified"/"Not disclosed", matching the field's existing value.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from collections import Counter

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
NEW_FILL = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
CLOSING_FILL = PatternFill(start_color="FED7AA", end_color="FED7AA", fill_type="solid")

COLUMNS = [
    ("company", "Company"), ("job_title", "Job Title"), ("job_id", "Job ID"),
    ("location_raw", "Location"), ("work_mode", "Work Mode"),
    ("employment_type", "Employment Type"), ("experience_raw", "Experience"),
    ("education_required", "Degree"), ("branch_required", "Branch"),
    ("salary_raw", "Salary / CTC"), ("currency", "Currency"),
    ("date_posted", "Date Posted"), ("application_deadline", "Deadline"),
    ("deadline_status", "Deadline Status"), ("apply_url", "Apply Link"),
    ("job_url", "Job URL"), ("fresher_eligible", "Fresher Eligible"),
    ("internship", "Internship"), ("cse_relevant", "CSE Relevant"),
    ("source_website", "Source"), ("status", "Status"),
    ("scraped_at", "Scraped At"),
]


def _write_sheet(ws, jobs, title):
    ws.title = title
    for col_idx, (_, header) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"

    for row_idx, job in enumerate(jobs, start=2):
        for col_idx, (key, _) in enumerate(COLUMNS, start=1):
            value = job.get(key)
            if isinstance(value, list):
                value = ", ".join(value) if value else "Not specified"
            if value is None:
                value = "Not specified"
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if key == "apply_url" and value not in ("Not specified", ""):
                cell.hyperlink = value
                cell.font = Font(color="2563EB", underline="single")
            if key == "status" and value == "NEW":
                cell.fill = NEW_FILL
            if key == "deadline_status" and value == "Closing Soon":
                cell.fill = CLOSING_FILL

    if jobs:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(jobs)+1}"
    for col_idx, (key, header) in enumerate(COLUMNS, start=1):
        width = max(len(header), 14)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _write_stats_sheet(ws, jobs):
    ws.title = "Statistics"
    stats = [
        ("Total Jobs", len(jobs)),
        ("New Jobs", sum(1 for j in jobs if j.get("status") == "NEW")),
        ("Companies Found", len(set(j.get("company") for j in jobs))),
        ("Fresher Jobs", sum(1 for j in jobs if j.get("fresher_eligible"))),
        ("Internships", sum(1 for j in jobs if j.get("internship"))),
        ("CSE Jobs", sum(1 for j in jobs if j.get("cse_relevant"))),
        ("Jobs With Salary Disclosed", sum(1 for j in jobs if j.get("salary_raw") not in (None, "Not disclosed"))),
        ("Jobs With Deadline", sum(1 for j in jobs if j.get("application_deadline"))),
    ]
    ws.cell(row=1, column=1, value="Metric").font = HEADER_FONT
    ws.cell(row=1, column=2, value="Value").font = HEADER_FONT
    ws.cell(row=1, column=1).fill = HEADER_FILL
    ws.cell(row=1, column=2).fill = HEADER_FILL
    for i, (label, value) in enumerate(stats, start=2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 15

    row = len(stats) + 3
    ws.cell(row=row, column=1, value="Jobs By Company").font = HEADER_FONT
    row += 1
    for company, count in Counter(j.get("company") for j in jobs).most_common():
        ws.cell(row=row, column=1, value=company)
        ws.cell(row=row, column=2, value=count)
        row += 1


def generate_xlsx(jobs: list, output_path: str):
    if not jobs:
        raise ValueError("generate_xlsx called with an empty job list -- refusing to write an empty workbook silently")

    wb = Workbook()
    _write_sheet(wb.active, jobs, "All Jobs")

    new_jobs = [j for j in jobs if j.get("status") == "NEW"]
    _write_sheet(wb.create_sheet(), new_jobs, "New Jobs")

    fresher_jobs = [j for j in jobs if j.get("fresher_eligible")]
    _write_sheet(wb.create_sheet(), fresher_jobs, "Fresher Jobs")

    internships = [j for j in jobs if j.get("internship")]
    _write_sheet(wb.create_sheet(), internships, "Internships")

    cse_jobs = [j for j in jobs if j.get("cse_relevant")]
    _write_sheet(wb.create_sheet(), cse_jobs, "CSE Jobs")

    closing_soon = [j for j in jobs if j.get("deadline_status") == "Closing Soon"]
    _write_sheet(wb.create_sheet(), closing_soon, "Closing Soon")

    _write_stats_sheet(wb.create_sheet(), jobs)

    wb.save(output_path)
    return {
        "total": len(jobs), "new": len(new_jobs), "fresher": len(fresher_jobs),
        "internships": len(internships), "cse": len(cse_jobs), "closing_soon": len(closing_soon),
    }


if __name__ == "__main__":
    # Self-test with synthetic data -- verifies the file is actually
    # produced, has the right sheets, and the counts match reality.
    import os
    import tempfile

    sample_jobs = [
        {"company": "Stripe", "job_title": "Software Engineer", "job_id": "1", "location_raw": "SF",
         "cse_relevant": True, "fresher_eligible": None, "internship": None, "status": "NEW",
         "apply_url": "https://stripe.com/jobs/1", "deadline_status": "Open"},
        {"company": "Stripe", "job_title": "SDE Intern", "job_id": "2", "location_raw": "NYC",
         "cse_relevant": True, "fresher_eligible": True, "internship": True, "status": "UNCHANGED",
         "apply_url": "https://stripe.com/jobs/2", "deadline_status": "Closing Soon"},
        {"company": "Anthropic", "job_title": "Sales Manager", "job_id": "3", "location_raw": "Remote",
         "cse_relevant": False, "fresher_eligible": None, "internship": None, "status": "NEW",
         "apply_url": "https://anthropic.com/jobs/3", "deadline_status": "Deadline Not Specified"},
    ]

    out_path = os.path.join(tempfile.gettempdir(), "test_jobs.xlsx")
    result = generate_xlsx(sample_jobs, out_path)
    print("generate_xlsx returned:", result)

    assert os.path.exists(out_path), "FAIL: file was not created"
    from openpyxl import load_workbook
    wb = load_workbook(out_path)
    expected_sheets = {"All Jobs", "New Jobs", "Fresher Jobs", "Internships", "CSE Jobs", "Closing Soon", "Statistics"}
    actual_sheets = set(wb.sheetnames)
    assert actual_sheets == expected_sheets, f"FAIL: sheet mismatch. got {actual_sheets}"
    assert wb["All Jobs"].max_row == 4, f"FAIL: expected 4 rows (1 header + 3 jobs), got {wb['All Jobs'].max_row}"
    assert wb["CSE Jobs"].max_row == 3, f"FAIL: expected 3 rows (1 header + 2 CSE jobs), got {wb['CSE Jobs'].max_row}"
    assert wb["Internships"].max_row == 2, f"FAIL: expected 2 rows (1 header + 1 internship), got {wb['Internships'].max_row}"
    assert result["total"] == 3 and result["cse"] == 2 and result["internships"] == 1

    print("ALL SELF-TESTS PASSED")
    os.remove(out_path)
