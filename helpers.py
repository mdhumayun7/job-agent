import json, re, hashlib, sqlite3, time, logging
from pathlib import Path
from datetime import datetime, timedelta
from difflib import SequenceMatcher
from collections import OrderedDict

# Load .env
_env = Path(".env")
if _env.exists():
    for line in _env.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k,v = line.split("=",1)
            import os; os.environ.setdefault(k.strip(), v.strip())

import os
def _csv(name, default):
    v = os.getenv(name)
    return [x.strip() for x in v.split(",")] if v else default

YOUR_SKILLS    = _csv("YOUR_SKILLS", ["python","machine learning","flask","sql","aws","git"])
PREFERRED_CITIES = _csv("CITIES", ["Bangalore","Hyderabad","Pune","Mumbai","Remote"])
ALLOWED_JOB_TYPES = _csv("ALLOWED_JOB_TYPES", ["Full Time","Internship"])
NEGATIVE_KEYWORDS = _csv("NEGATIVE_KEYWORDS", ["sales","bpo","telecalling","marketing","hr recruiter"])
MIN_MATCH_SCORE = int(os.getenv("MIN_MATCH_SCORE","5"))
REMOTE_ONLY = os.getenv("REMOTE_ONLY","false").lower() == "true"
MAX_EXCEL_JOBS_PER_SOURCE = int(os.getenv("MAX_EXCEL_JOBS_PER_SOURCE","100"))

OUTPUT_FILE  = os.getenv("OUTPUT_FILE","output/jobs_found.json")
EXCEL_FILE   = os.getenv("EXCEL_FILE","output/jobs_found.xlsx")
TRACKER_FILE = os.getenv("TRACKER_FILE","output/apply_tracker.xlsx")
DB_FILE      = os.getenv("DB_FILE","data/jobs.db")
STATUS_COLORS = {"Saved":"E2E8F0","Applied":"BFDBFE","Interview":"86EFAC","Rejected":"FCA5A5","Offer":"4ADE80","Follow Up":"FDE68A"}

try:
    from config import OUTPUT_FILE, EXCEL_FILE, TRACKER_FILE, DB_FILE, STATUS_COLORS
    from config import PREFERRED_CITIES as PC
    PREFERRED_CITIES = PC
except: pass

LOG_DIR = Path("logs")
LOG_DIR.mkdir(exist_ok=True)
logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.FileHandler(LOG_DIR/"scraper.log",encoding="utf-8"), logging.StreamHandler()])
log = logging.getLogger(__name__)

JOBS_FILE = OUTPUT_FILE

# ── Scoring ───────────────────────────────────────────────────────────────────
TITLE_BOOST = ["fresher","junior","entry","trainee","associate","graduate","intern",
               "0-1","0-2","analyst","developer","engineer","scientist","designer","sde"]
NEGATIVE_EXP = ["senior","lead","manager","director","head","principal","architect",
                "5+ years","7+ years","10+ years","vp ","president","cto","ceo",
                "3+ years","4+ years","15+ years"]

def score_job(job):
    score = 0
    title = (job.get("title") or "").lower()
    desc  = (job.get("description") or "").lower()
    sal   = (job.get("salary") or "").lower()
    src   = (job.get("source_group") or job.get("source","")).lower()
    loc   = (job.get("location") or "").lower()
    text  = f"{title} {desc}"

    # Negative keyword filter — return 0 if irrelevant job
    neg_kws = [k.lower() for k in NEGATIVE_KEYWORDS]
    neg_in_title = sum(1 for k in neg_kws if k in title)
    if neg_in_title >= 2: return 0
    if neg_in_title == 1 and not any(s.lower() in title for s in YOUR_SKILLS): return 0

    # Skill scoring
    for skill in YOUR_SKILLS:
        sl = skill.lower()
        if sl in title: score += 18
        elif sl in desc: score += 8

    # Fresher keywords in title
    for kw in TITLE_BOOST:
        if kw in title: score += 8; break

    # Experience penalty
    for kw in NEGATIVE_EXP:
        if kw in text: score -= 20; break

    # Salary boost
    if any(x in sal for x in ["lpa","lakh","stipend","per annum"]):
        nums = re.findall(r"\d+\.?\d*", sal.replace(",",""))
        if nums:
            mx = max(float(n) for n in nums)
            if mx > 200000: mx /= 100000
            if mx >= 5: score += 10
            elif mx >= 3: score += 5

    # Source boost
    if "naukri" in src: score += 5
    if "linkedin" in src: score += 5
    if "internshala" in src: score += 3

    # Location preference
    pref = [c.lower() for c in PREFERRED_CITIES]
    if any(c in loc for c in pref): score += 5
    if "remote" in loc: score += 3

    # Freshness boost
    posted = job.get("posted_date","")
    if posted:
        try:
            if "day" in str(posted).lower():
                days = int(re.search(r"\d+", str(posted)).group())
                if days <= 1: score += 8
                elif days <= 3: score += 5
                elif days <= 7: score += 2
            elif "hour" in str(posted).lower() or "minute" in str(posted).lower():
                score += 10
            elif len(str(posted)) >= 10:
                post_date = datetime.fromisoformat(str(posted)[:10])
                days_old = (datetime.now() - post_date).days
                if days_old <= 1: score += 8
                elif days_old <= 3: score += 5
                elif days_old <= 7: score += 2
                elif days_old > 30: score -= 5
        except: pass

    return max(0, min(score, 100))

def normalize_source_name(source):
    s = (source or "").strip().lower()
    if "linkedin" in s: return "LinkedIn"
    if "indeed" in s: return "Indeed"
    if "naukri" in s: return "Naukri"
    if "angel" in s or "wellfound" in s: return "Wellfound"
    if "unstop" in s: return "Unstop"
    if "foundit" in s or "monster" in s: return "Foundit"
    if "internshala" in s: return "Internshala"
    if "shine" in s: return "Shine"
    if "timesjobs" in s or "times" in s: return "TimesJobs"
    if "glassdoor" in s: return "Glassdoor"
    if "hacker" in s: return "HackerEarth"
    if "govt" in s or "sarkari" in s: return "Government"
    return source or "Other"

def infer_job_type(job):
    text = f"{job.get('title','')} {job.get('description','')}".lower()
    if any(x in text for x in ["intern","internship","trainee","apprentice"]): return "Internship"
    return "Full Time"

def make_job_id(job):
    key = f"{job.get('title','')}-{job.get('company','')}-{normalize_source_name(job.get('source',''))}".lower()
    return hashlib.md5(key.encode()).hexdigest()[:12]

def deduplicate(jobs):
    unique, seen_urls, seen_keys = [], set(), set()
    for job in jobs:
        job["source_group"] = normalize_source_name(job.get("source",""))
        job["job_type"]     = infer_job_type(job)
        job["job_id"]       = job.get("job_id") or make_job_id(job)
        url = (job.get("apply_url") or "").strip()
        if url and url in seen_urls: continue
        # Title+Company exact dedup
        key = f"{job.get('title','').lower()[:40]}-{job.get('company','').lower()[:30]}"
        if key in seen_keys: continue
        # Fuzzy dedup — check last 100 jobs
        dup = False
        for ex in unique[-100:]:
            t_sim = SequenceMatcher(None, job.get("title","").lower(), ex.get("title","").lower()).ratio()
            c_sim = SequenceMatcher(None, job.get("company","").lower(), ex.get("company","").lower()).ratio()
            if t_sim > 0.88 and c_sim > 0.88:
                if job.get("match_score",0) > ex.get("match_score",0):
                    ex.update(job)
                dup = True
                break
        if not dup:
            unique.append(job)
            if url: seen_urls.add(url)
            seen_keys.add(key)
    log.info(f"Dedup: {len(jobs)} -> {len(unique)} unique")
    return unique

def apply_filters(jobs):
    neg_kws = [k.lower() for k in NEGATIVE_KEYWORDS]
    out = []
    for job in jobs:
        score = job.get("match_score", 0)
        if score < MIN_MATCH_SCORE: continue
        jtype = job.get("job_type","Full Time").lower()
        allowed = [t.lower() for t in ALLOWED_JOB_TYPES]
        if allowed and jtype not in allowed: continue
        # Hard negative filter on title
        title = job.get("title","").lower()
        neg_count = sum(1 for k in neg_kws if k in title)
        if neg_count >= 2: continue
        if neg_count == 1 and score < 15: continue
        out.append(job)
    log.info(f"Filter: {len(jobs)} -> {len(out)} jobs")
    return out

def init_database():
    Path(DB_FILE).parent.mkdir(exist_ok=True)
    with sqlite3.connect(DB_FILE) as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS jobs (
            job_id TEXT PRIMARY KEY, title TEXT, company TEXT, location TEXT,
            source TEXT, source_group TEXT, salary TEXT, apply_url TEXT,
            posted_date TEXT, match_score INTEGER, job_type TEXT,
            status TEXT DEFAULT 'Saved', notes TEXT DEFAULT '',
            applied_date TEXT DEFAULT '', interview_date TEXT DEFAULT '',
            follow_up_date TEXT DEFAULT '', scraped_at TEXT, raw_json TEXT)""")
        conn.commit()

def save_jobs_to_db(jobs):
    init_database()
    with sqlite3.connect(DB_FILE) as conn:
        for job in jobs:
            conn.execute("""INSERT INTO jobs
                (job_id,title,company,location,source,source_group,salary,apply_url,
                posted_date,match_score,job_type,status,scraped_at,raw_json)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                ON CONFLICT(job_id) DO UPDATE SET
                title=excluded.title, company=excluded.company,
                location=excluded.location, salary=excluded.salary,
                match_score=excluded.match_score, scraped_at=excluded.scraped_at,
                raw_json=excluded.raw_json""",
                (job.get("job_id"), job.get("title",""), job.get("company",""),
                 job.get("location",""), job.get("source",""), job.get("source_group",""),
                 job.get("salary",""), job.get("apply_url",""), job.get("posted_date",""),
                 job.get("match_score",0), job.get("job_type",""),
                 job.get("status","Saved"), job.get("scraped_at",""),
                 json.dumps(job, ensure_ascii=False)))
        conn.commit()
    log.info(f"DB: saved {len(jobs)} jobs")

def export_excel(jobs):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "All Jobs"
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "D2"
        hdrs = ["#","Score","Grade","Title","Company","Location","Salary","Type","Source","Status","Posted","Matched Skills","Apply Link"]
        widths = [4,7,13,36,22,18,16,12,14,10,11,28,45]
        for col,(h,w) in enumerate(zip(hdrs,widths),1):
            c = ws.cell(row=1,column=col,value=h)
            c.fill = PatternFill("solid",fgColor="4F46E5")
            c.font = Font(color="FFFFFF",bold=True,size=10)
            c.alignment = Alignment(horizontal="center",vertical="center")
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.row_dimensions[1].height = 28
        def grade(s): return "Excellent" if s>=60 else ("Good" if s>=40 else ("Fair" if s>=20 else "Low"))
        def sfill(s):
            if s>=60: return "14532D","86EFAC"
            elif s>=40: return "1E3A5F","60A5FA"
            elif s>=20: return "422006","FDE68A"
            return "450A0A","FCA5A5"
        for row,job in enumerate(jobs,2):
            sc = job.get("match_score",0)
            rf = PatternFill("solid",fgColor="0F172A" if row%2==0 else "1E293B")
            matched = ", ".join(job.get("matched_skills",[])[:5]) if isinstance(job.get("matched_skills",[]),list) else str(job.get("matched_skills",""))
            vals=[row-1,sc,grade(sc),job.get("title",""),job.get("company",""),
                  job.get("location",""),job.get("salary",""),job.get("job_type",""),
                  job.get("source_group",""),job.get("status","Saved"),
                  job.get("posted_date",""),matched,job.get("apply_url","")]
            for col,val in enumerate(vals,1):
                c = ws.cell(row=row,column=col,value=val)
                c.fill = rf
                c.alignment = Alignment(horizontal="center" if col in [1,2,3,8,9,10,11] else "left",
                                       vertical="center",wrap_text=(col in [4,12]))
                c.font = Font(size=9,color="E2E8F0")
            bg,fg = sfill(sc)
            sc_c = ws.cell(row=row,column=2)
            sc_c.fill = PatternFill("solid",fgColor=bg)
            sc_c.font = Font(bold=True,size=10,color=fg)
            sc_c.alignment = Alignment(horizontal="center",vertical="center")
            ws.row_dimensions[row].height = 18
        Path(EXCEL_FILE).parent.mkdir(exist_ok=True)
        wb.save(EXCEL_FILE)
        log.info(f"Excel saved: {EXCEL_FILE}")
    except Exception as e:
        log.error(f"Excel error: {e}")

def export_tracker(jobs):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Apply Tracker"
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "C2"
        hdrs = ["#","Score","Title","Company","Location","Salary","Source","Status",
                "Applied Date","Interview Date","Follow Up","Notes","Apply Link"]
        widths = [4,7,32,22,18,16,14,12,14,14,14,25,42]
        for col,(h,w) in enumerate(zip(hdrs,widths),1):
            c = ws.cell(row=1,column=col,value=h)
            c.fill = PatternFill("solid",fgColor="1E293B")
            c.font = Font(color="FFFFFF",bold=True,size=10)
            c.alignment = Alignment(horizontal="center",vertical="center")
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.row_dimensions[1].height = 28
        st_bg = {"Saved":"1E293B","Applied":"1E3A5F","Interview":"064E3B","Rejected":"450A0A","Offer":"052E16","Follow Up":"422006"}
        st_fg = {"Saved":"94A3B8","Applied":"60A5FA","Interview":"34D399","Rejected":"F87171","Offer":"4ADE80","Follow Up":"FDE68A"}
        for row,job in enumerate(jobs,2):
            sc = job.get("match_score",0)
            status = job.get("status","Saved")
            rf = PatternFill("solid",fgColor="0F172A" if row%2==0 else "1E293B")
            vals=[row-1,sc,job.get("title",""),job.get("company",""),job.get("location",""),
                  job.get("salary",""),job.get("source_group",""),status,
                  job.get("applied_date",""),job.get("interview_date",""),
                  job.get("follow_up_date",""),job.get("notes",""),job.get("apply_url","")]
            for col,val in enumerate(vals,1):
                c = ws.cell(row=row,column=col,value=val)
                c.fill = rf
                c.alignment = Alignment(horizontal="center" if col in [1,2,8,9,10,11] else "left",
                                       vertical="center",wrap_text=(col in [3,12]))
                c.font = Font(size=9,color="E2E8F0")
            st_c = ws.cell(row=row,column=8)
            st_c.fill = PatternFill("solid",fgColor=st_bg.get(status,"1E293B"))
            st_c.font = Font(color=st_fg.get(status,"94A3B8"),bold=True,size=9)
            st_c.alignment = Alignment(horizontal="center",vertical="center")
            ws.row_dimensions[row].height = 18
        Path(TRACKER_FILE).parent.mkdir(exist_ok=True)
        wb.save(TRACKER_FILE)
        log.info(f"Tracker saved: {TRACKER_FILE}")
    except Exception as e:
        log.error(f"Tracker error: {e}")

def save_results(jobs):
    Path("output").mkdir(exist_ok=True)
    for job in jobs:
        job["match_score"]  = score_job(job)
        job["scraped_at"]   = datetime.now().isoformat()
        job["source_group"] = normalize_source_name(job.get("source",""))
        job["job_type"]     = infer_job_type(job)
        job["status"]       = job.get("status","Saved")
        job["job_id"]       = job.get("job_id") or make_job_id(job)
    jobs = deduplicate(jobs)
    jobs = apply_filters(jobs)
    jobs.sort(key=lambda j: j["match_score"], reverse=True)
    with open(OUTPUT_FILE,"w",encoding="utf-8") as f:
        json.dump(jobs,f,ensure_ascii=False,indent=2)
    try:
        from config import JOBS_FILE as JF
        with open(JF,"w",encoding="utf-8") as f:
            json.dump(jobs,f,ensure_ascii=False,indent=2)
    except: pass
    log.info(f"Saved {len(jobs)} jobs")
    save_jobs_to_db(jobs)
    export_excel(jobs)
    export_tracker(jobs)
    return jobs

def polite_wait(seconds=2):
    time.sleep(seconds)
