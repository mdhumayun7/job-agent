import json
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from config import ENRICHED_OUTPUT_FILE, JOBS_FILE, YOUR_SKILLS, PERSONAL_INFO, PROFILE_SUMMARY

TECH_QUESTIONS = {
    "python":         ["Explain Python GIL","List vs Tuple difference","What are decorators?","Explain generators"],
    "machine learning":["Bias-variance tradeoff","Overfitting solutions","Explain gradient descent","Cross-validation"],
    "sql":            ["Joins explained","GROUP BY vs HAVING","Indexes and performance","Window functions"],
    "react":          ["Virtual DOM","useState vs useReducer","useEffect cleanup","Props vs State"],
    "flask":          ["Flask vs Django","Blueprints","Request context","REST API design"],
    "deep learning":  ["CNN vs RNN","Backpropagation","Batch normalization","Dropout regularization"],
    "data analysis":  ["Pandas groupby","Handling missing data","Data normalization","EDA steps"],
    "git":            ["Git rebase vs merge","Cherry-pick","Git stash","Branch strategies"],
    "linux":          ["File permissions","Process management","Cron jobs","Shell scripting basics"],
    "aws":            ["EC2 vs Lambda","S3 storage classes","IAM roles","VPC basics"],
}
HR_QUESTIONS = [
    "Tell me about yourself",
    "Why do you want to join this company?",
    "What are your strengths and weaknesses?",
    "Where do you see yourself in 5 years?",
    "Describe a challenging project you worked on",
    "Why should we hire you?",
]
CODING_PROBLEMS = [
    "Reverse a string / linked list",
    "Find duplicates in array",
    "Binary search implementation",
    "Fibonacci with memoization",
    "Valid parentheses (stack)",
    "Two sum problem",
]

def get_prep(job):
    text    = f"{job.get('title','')} {job.get('description','')}".lower()
    matched = job.get("matched_skills",[]) or []
    topics  = job.get("interview_focus_topics",[]) or matched[:4] or ["DSA","OOPs","SQL","System Design"]
    tech_qs = []
    for skill in matched[:4]:
        qs = TECH_QUESTIONS.get(skill.lower(),[])
        tech_qs.extend(qs[:2])
    if not tech_qs:
        tech_qs = ["Explain your strongest technical skill","Walk me through a project you built","How do you debug complex issues?","Explain OOP concepts with example"]
    score = job.get("ai_match_score",job.get("match_score",0))
    diff  = job.get("interview_difficulty","Medium")
    prep_days = 3 if score>=60 else (5 if score>=30 else 7)
    tip = f"Focus on {topics[0] if topics else 'core CS fundamentals'} — most likely to be asked for this role."
    return {
        "technical_questions":    "\n".join([f"{i+1}. {q}" for i,q in enumerate(tech_qs[:6])]),
        "hr_questions":           "\n".join([f"{i+1}. {q}" for i,q in enumerate(HR_QUESTIONS[:4])]),
        "coding_problems":        "\n".join([f"• {p}" for p in CODING_PROBLEMS[:4]]),
        "topics_to_study":        "\n".join([f"• {t}" for t in topics[:5]]),
        "prep_days_recommended":  prep_days,
        "difficulty":             diff,
        "tips":                   tip,
    }

def save_excel(jobs, path="output/interview_prep.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Interview Prep"
    headers = ["Score","Title","Company","Source","Difficulty","Prep Days","Technical Questions","HR Questions","Coding Problems","Topics to Study","Tips","Apply Link"]
    hfill = PatternFill("solid",fgColor="1E293B")
    hfont = Font(color="FFFFFF",bold=True)
    for col,h in enumerate(headers,1):
        c = ws.cell(row=1,column=col,value=h)
        c.fill=hfill; c.font=hfont
        c.alignment=Alignment(horizontal="center")
    diff_colors = {"Easy":"86EFAC","Medium":"FDE68A","Hard":"FCA5A5"}
    for row,job in enumerate(jobs,2):
        sc = job.get("ai_match_score",job.get("match_score",0))
        color = "86EFAC" if sc>=60 else ("FDE68A" if sc>=30 else "FCA5A5")
        ws.cell(row=row,column=1,value=sc).fill=PatternFill("solid",fgColor=color)
        ws.cell(row=row,column=2,value=job.get("title",""))
        ws.cell(row=row,column=3,value=job.get("company",""))
        ws.cell(row=row,column=4,value=job.get("source_group",job.get("source","")))
        diff=job.get("difficulty","Medium")
        dc=ws.cell(row=row,column=5,value=diff)
        dc.fill=PatternFill("solid",fgColor=diff_colors.get(diff,"E2E8F0"))
        ws.cell(row=row,column=6,value=job.get("prep_days_recommended",5))
        ws.cell(row=row,column=7,value=job.get("technical_questions","")).alignment=Alignment(wrap_text=True)
        ws.cell(row=row,column=8,value=job.get("hr_questions","")).alignment=Alignment(wrap_text=True)
        ws.cell(row=row,column=9,value=job.get("coding_problems","")).alignment=Alignment(wrap_text=True)
        ws.cell(row=row,column=10,value=job.get("topics_to_study","")).alignment=Alignment(wrap_text=True)
        ws.cell(row=row,column=11,value=job.get("tips","")).alignment=Alignment(wrap_text=True)
        ws.cell(row=row,column=12,value=job.get("apply_url",""))
    for col in ws.columns:
        mx=max((len(str(c.value or "")) for c in col),default=10)
        ws.column_dimensions[col[0].column_letter].width=min(mx+4,60)
    Path(path).parent.mkdir(exist_ok=True)
    wb.save(path)
    print(f"  Saved: {path}")

def main():
    src = Path(ENRICHED_OUTPUT_FILE) if Path(ENRICHED_OUTPUT_FILE).exists() else Path(JOBS_FILE)
    with open(src,"r",encoding="utf-8") as f: jobs=json.load(f)
    top = sorted(jobs,key=lambda j:j.get("ai_match_score",j.get("match_score",0)),reverse=True)[:50]
    print(f"\nGenerating interview prep for top {len(top)} jobs...")
    prepped=[]
    for i,job in enumerate(top,1):
        p=get_prep(job)
        prepped.append({**job,**p})
        if i%10==0: print(f"  {i}/{len(top)} done")
    save_excel(prepped)
    print(f"Done! {len(prepped)} jobs prepared.")

if __name__ == "__main__":
    main()
