import json, re
from pathlib import Path
from collections import Counter, OrderedDict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config import EXCEL_FILE, TRACKER_FILE, ENRICHED_OUTPUT_FILE, JOBS_FILE, YOUR_SKILLS, PERSONAL_INFO, PROFILE_SUMMARY

# ── Styles ────────────────────────────────────────────────────────────────────
def border():
    s = Side(style="thin", color="1E293B")
    return Border(left=s, right=s, top=s, bottom=s)

def hdr(cell, bg="4F46E5", fg="FFFFFF", size=10):
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.font = Font(color=fg, bold=True, size=size)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border()

def dat(cell, wrap=False, align="left", size=9):
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    cell.border = border()
    cell.font = Font(size=size, color="E2E8F0")

def score_style(cell, score):
    if score >= 60:   cell.fill, fgc = PatternFill("solid", fgColor="14532D"), "86EFAC"
    elif score >= 40: cell.fill, fgc = PatternFill("solid", fgColor="1E3A5F"), "60A5FA"
    elif score >= 20: cell.fill, fgc = PatternFill("solid", fgColor="422006"), "FDE68A"
    else:             cell.fill, fgc = PatternFill("solid", fgColor="450A0A"), "FCA5A5"
    cell.font = Font(bold=True, size=10, color=fgc)
    cell.alignment = Alignment(horizontal="center", vertical="center")

def grade(score):
    if score >= 60: return "🟢 Excellent"
    elif score >= 40: return "🔵 Good"
    elif score >= 20: return "🟡 Fair"
    return "🔴 Low"

def row_fill(row_num):
    return PatternFill("solid", fgColor="0F172A" if row_num % 2 == 0 else "1E293B")

def load_jobs():
    src = Path(ENRICHED_OUTPUT_FILE) if Path(ENRICHED_OUTPUT_FILE).exists() else Path(JOBS_FILE)
    with open(src, "r", encoding="utf-8") as f:
        return json.load(f)

# ── Main Jobs Excel ───────────────────────────────────────────────────────────
def build_jobs_excel(jobs):
    wb = openpyxl.Workbook()

    # Sheet 1: All Jobs
    ws = wb.active
    ws.title = "All Jobs"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "D2"

    hdrs = ["#", "Score", "Grade", "Title", "Company", "Location", "Salary", "Type", "Source", "Status", "Posted", "Matched Skills", "Apply Link"]
    widths = [4, 7, 13, 36, 22, 18, 16, 12, 14, 10, 11, 28, 45]

    for col, (h, w) in enumerate(zip(hdrs, widths), 1):
        c = ws.cell(row=1, column=col, value=h)
        hdr(c)
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 30

    for row, job in enumerate(jobs, 2):
        sc = job.get("ai_match_score", job.get("match_score", 0))
        rf = row_fill(row)
        matched = ", ".join(job.get("matched_skills", [])[:5])
        vals = [row-1, sc, grade(sc), job.get("title",""), job.get("company",""),
                job.get("location",""), job.get("salary","Not mentioned"),
                job.get("job_type","Full Time"), job.get("source_group", job.get("source","")),
                job.get("status","Saved"), job.get("posted_date",""), matched, job.get("apply_url","")]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=safe_str(val))
            c.fill = rf
            dat(c, wrap=(col in [4,12]), align="center" if col in [1,2,3,8,9,10,11] else "left")
        score_style(ws.cell(row=row, column=2), sc)
        # Status color
        st_cell = ws.cell(row=row, column=10)
        st_colors = {"Saved":"334155","Applied":"1E3A5F","Interview":"064E3B","Rejected":"450A0A","Offer":"052E16"}
        st_fg = {"Saved":"94A3B8","Applied":"60A5FA","Interview":"34D399","Rejected":"F87171","Offer":"4ADE80"}
        status = job.get("status","Saved")
        st_cell.fill = PatternFill("solid", fgColor=st_colors.get(status,"334155"))
        st_cell.font = Font(color=st_fg.get(status,"94A3B8"), bold=True, size=9)
        st_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 18

    # Sheet 2: Stats
    ws2 = wb.create_sheet("📊 Stats")
    ws2.sheet_view.showGridLines = False
    ws2.merge_cells("A1:G1")
    tc = ws2.cell(row=1, column=1, value="📊 Job Search Analytics Dashboard")
    tc.font = Font(size=18, bold=True, color="818CF8")
    tc.alignment = Alignment(horizontal="center", vertical="center")
    tc.fill = PatternFill("solid", fgColor="0F172A")
    ws2.row_dimensions[1].height = 40

    # Source stats
    ws2.cell(row=3, column=1, value="🌐 Platform").fill = PatternFill("solid", fgColor="4F46E5")
    ws2.cell(row=3, column=1).font = Font(color="FFFFFF", bold=True)
    ws2.cell(row=3, column=2, value="Jobs Found").fill = PatternFill("solid", fgColor="4F46E5")
    ws2.cell(row=3, column=2).font = Font(color="FFFFFF", bold=True)
    sources = Counter(j.get("source_group", j.get("source","Other")) for j in jobs)
    for i, (src, cnt) in enumerate(sources.most_common(), 4):
        c1 = ws2.cell(row=i, column=1, value=src)
        c2 = ws2.cell(row=i, column=2, value=cnt)
        c1.fill = c2.fill = PatternFill("solid", fgColor="1E293B")
        c1.font = Font(color="E2E8F0", size=10)
        c2.font = Font(color="818CF8", bold=True, size=11)
        c1.alignment = c2.alignment = Alignment(horizontal="left", vertical="center")

    # Score ranges
    ws2.cell(row=3, column=4, value="📊 Score Range").fill = PatternFill("solid", fgColor="7C3AED")
    ws2.cell(row=3, column=4).font = Font(color="FFFFFF", bold=True)
    ws2.cell(row=3, column=5, value="Count").fill = PatternFill("solid", fgColor="7C3AED")
    ws2.cell(row=3, column=5).font = Font(color="FFFFFF", bold=True)
    ranges = [("🟢 60-100 Excellent",60,101),("🔵 40-59 Good",40,60),("🟡 20-39 Fair",20,40),("🔴 0-19 Low",0,20)]
    for i, (lbl,lo,hi) in enumerate(ranges, 4):
        cnt = sum(1 for j in jobs if lo <= j.get("ai_match_score",j.get("match_score",0)) < hi)
        c1 = ws2.cell(row=i, column=4, value=lbl)
        c2 = ws2.cell(row=i, column=5, value=cnt)
        c1.fill = c2.fill = PatternFill("solid", fgColor="1E293B")
        c1.font = Font(color="E2E8F0", size=10)
        c2.font = Font(color="A78BFA", bold=True, size=11)

    # Top skills
    sk_row = 3 + len(sources) + 3
    ws2.cell(row=sk_row, column=1, value="🧠 Top Skills In Demand").fill = PatternFill("solid", fgColor="0EA5E9")
    ws2.cell(row=sk_row, column=1).font = Font(color="FFFFFF", bold=True)
    ws2.cell(row=sk_row, column=2, value="Jobs Count").fill = PatternFill("solid", fgColor="0EA5E9")
    ws2.cell(row=sk_row, column=2).font = Font(color="FFFFFF", bold=True)
    sk_cnt = Counter()
    for j in jobs:
        for sk in j.get("matched_skills",[]): sk_cnt[sk] += 1
    for i, (sk, cnt) in enumerate(sk_cnt.most_common(12), sk_row+1):
        c1 = ws2.cell(row=i, column=1, value=sk)
        c2 = ws2.cell(row=i, column=2, value=cnt)
        c1.fill = c2.fill = PatternFill("solid", fgColor="1E293B")
        c1.font = Font(color="E2E8F0", size=10)
        c2.font = Font(color="38BDF8", bold=True, size=11)

    # Top companies
    ws2.cell(row=sk_row, column=4, value="🏢 Top Companies").fill = PatternFill("solid", fgColor="059669")
    ws2.cell(row=sk_row, column=4).font = Font(color="FFFFFF", bold=True)
    ws2.cell(row=sk_row, column=5, value="Jobs").fill = PatternFill("solid", fgColor="059669")
    ws2.cell(row=sk_row, column=5).font = Font(color="FFFFFF", bold=True)
    comp_cnt = Counter(j.get("company","") for j in jobs if j.get("company",""))
    for i, (comp, cnt) in enumerate(comp_cnt.most_common(12), sk_row+1):
        c1 = ws2.cell(row=i, column=4, value=comp)
        c2 = ws2.cell(row=i, column=5, value=cnt)
        c1.fill = c2.fill = PatternFill("solid", fgColor="1E293B")
        c1.font = Font(color="E2E8F0", size=10)
        c2.font = Font(color="34D399", bold=True, size=11)

    for col in ["A","B","C","D","E","F","G"]:
        ws2.column_dimensions[col].width = 30

    # Per-source sheets
    grouped = OrderedDict()
    for j in jobs:
        src = j.get("source_group", j.get("source","Other"))
        grouped.setdefault(src,[]).append(j)
    for src, src_jobs in grouped.items():
        ws_src = wb.create_sheet(title=src[:28])
        ws_src.sheet_view.showGridLines = False
        ws_src.freeze_panes = "D2"
        for col, (h, w) in enumerate(zip(hdrs, widths), 1):
            c = ws_src.cell(row=1, column=col, value=h)
            hdr(c)
            ws_src.column_dimensions[get_column_letter(col)].width = w
        ws_src.row_dimensions[1].height = 30
        for row, job in enumerate(src_jobs[:100], 2):
            sc = job.get("ai_match_score", job.get("match_score", 0))
            rf = row_fill(row)
            matched = ", ".join(job.get("matched_skills", [])[:5])
            vals = [row-1, sc, grade(sc), job.get("title",""), job.get("company",""),
                    job.get("location",""), job.get("salary","Not mentioned"),
                    job.get("job_type","Full Time"), job.get("source_group",job.get("source","")),
                    job.get("status","Saved"), job.get("posted_date",""), matched, job.get("apply_url","")]
            for col, val in enumerate(vals, 1):
                c = ws_src.cell(row=row, column=col, value=val)
                c.fill = rf
                dat(c, wrap=(col in [4,12]), align="center" if col in [1,2,3,8,9,10,11] else "left")
            score_style(ws_src.cell(row=row, column=2), sc)
            ws_src.row_dimensions[row].height = 18

    Path(EXCEL_FILE).parent.mkdir(exist_ok=True)
    wb.save(EXCEL_FILE)
    print(f"✅ Jobs Excel saved: {EXCEL_FILE}")

# ── Tailored Jobs Excel ───────────────────────────────────────────────────────
def safe_str(val):
    if isinstance(val, list): return ", ".join(str(x) for x in val)
    if val is None: return ""
    return str(val)

def build_tailored_excel(jobs):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tailored Applications"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "C2"

    hdrs = ["Score","Grade","Title","Company","Location","Source","One-Liner Pitch","Resume Objective","Cover Letter","ATS Keywords","Missing Skills","Apply Link"]
    widths = [7,13,32,22,18,14,35,40,55,30,25,40]

    for col,(h,w) in enumerate(zip(hdrs,widths),1):
        c = ws.cell(row=1,column=col,value=h)
        hdr(c, bg="7C3AED")
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 30

    for row, job in enumerate(jobs, 2):
        sc = job.get("ai_match_score", job.get("match_score",0))
        rf = row_fill(row)
        vals = [sc, grade(sc), job.get("title",""), job.get("company",""), job.get("location",""),
                job.get("source_group",job.get("source","")), job.get("one_liner",""),
                job.get("objective",""), job.get("cover_letter",""), job.get("ats_keywords",""),
                job.get("missing_skills",""), job.get("apply_url","")]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=safe_str(val))
            c.fill = rf
            dat(c, wrap=(col in [7,8,9,10]), align="center" if col in [1,2,6] else "left")
        score_style(ws.cell(row=row,column=1), sc)
        ws.row_dimensions[row].height = 80 if row <= 11 else 50

    out = "output/tailored_jobs.xlsx"
    Path(out).parent.mkdir(exist_ok=True)
    wb.save(out)
    print(f"✅ Tailored Excel saved: {out}")

# ── Interview Prep Excel ──────────────────────────────────────────────────────
def build_interview_excel(jobs):
    try:
        with open("output/interview_prep.xlsx","rb") as f:
            src_wb = openpyxl.load_workbook(f)
        src_ws = src_wb.active
    except:
        print("interview_prep.xlsx not found — run interview_prep.py first")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Interview Prep"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "C2"

    hdrs = ["Score","Title","Company","Source","Difficulty","Prep Days","Technical Questions","HR Questions","Coding Problems","Topics to Study","Tips","Apply Link"]
    widths = [7,32,22,14,10,10,45,35,30,28,35,40]
    diff_colors = {"Easy":"14532D","Medium":"422006","Hard":"450A0A"}
    diff_fg = {"Easy":"86EFAC","Medium":"FDE68A","Hard":"FCA5A5"}

    for col,(h,w) in enumerate(zip(hdrs,widths),1):
        c = ws.cell(row=1,column=col,value=h)
        hdr(c, bg="1E293B")
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 30

    # Copy data from existing interview_prep.xlsx
    for src_row in src_ws.iter_rows(min_row=2, values_only=True):
        row = ws.max_row + 1
        sc = src_row[0] if src_row[0] else 0
        rf = row_fill(row)
        for col, val in enumerate(src_row, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.fill = rf
            dat(c, wrap=(col in [7,8,9,10,11]), align="center" if col in [1,4,5,6] else "left", size=9)
        score_style(ws.cell(row=row,column=1), sc)
        diff = str(src_row[4] if len(src_row)>4 else "Medium")
        dc = ws.cell(row=row,column=5)
        dc.fill = PatternFill("solid",fgColor=diff_colors.get(diff,"422006"))
        dc.font = Font(color=diff_fg.get(diff,"FDE68A"),bold=True,size=9)
        dc.alignment = Alignment(horizontal="center",vertical="center")
        ws.row_dimensions[row].height = 60

    out = "output/interview_prep_styled.xlsx"
    wb.save(out)
    print(f"✅ Interview Excel saved: {out}")

# ── Apply Tracker Excel ───────────────────────────────────────────────────────
def build_tracker_excel(jobs):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Apply Tracker"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "C2"

    hdrs = ["#","Score","Title","Company","Location","Salary","Source","Status","Applied Date","Interview Date","Follow Up","Notes","Apply Link"]
    widths = [4,7,32,22,18,16,14,12,14,14,14,25,42]
    st_bg = {"Saved":"1E293B","Applied":"1E3A5F","Interview":"064E3B","Rejected":"450A0A","Offer":"052E16","Follow Up":"422006"}
    st_fg = {"Saved":"94A3B8","Applied":"60A5FA","Interview":"34D399","Rejected":"F87171","Offer":"4ADE80","Follow Up":"FDE68A"}

    for col,(h,w) in enumerate(zip(hdrs,widths),1):
        c = ws.cell(row=1,column=col,value=h)
        hdr(c, bg="0F172A")
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 30

    for row, job in enumerate(jobs, 2):
        sc = job.get("ai_match_score",job.get("match_score",0))
        status = job.get("status","Saved")
        rf = row_fill(row)
        vals = [row-1,sc,job.get("title",""),job.get("company",""),job.get("location",""),
                job.get("salary",""),job.get("source_group",job.get("source","")),status,
                job.get("applied_date",""),job.get("interview_date",""),
                job.get("follow_up_date",""),job.get("notes",""),job.get("apply_url","")]
        for col,val in enumerate(vals,1):
            c = ws.cell(row=row,column=col,value=val)
            c.fill = rf
            dat(c,wrap=(col in [3,12]),align="center" if col in [1,2,8,9,10,11] else "left")
        score_style(ws.cell(row=row,column=2),sc)
        sc2 = ws.cell(row=row,column=8)
        sc2.fill = PatternFill("solid",fgColor=st_bg.get(status,"1E293B"))
        sc2.font = Font(color=st_fg.get(status,"94A3B8"),bold=True,size=9)
        sc2.alignment = Alignment(horizontal="center",vertical="center")
        ws.row_dimensions[row].height = 18

    Path(TRACKER_FILE).parent.mkdir(exist_ok=True)
    wb.save(TRACKER_FILE)
    print(f"✅ Tracker Excel saved: {TRACKER_FILE}")

if __name__ == "__main__":
    jobs = load_jobs()
    jobs_sorted = sorted(jobs,key=lambda j:j.get("ai_match_score",j.get("match_score",0)),reverse=True)
    print(f"Processing {len(jobs_sorted)} jobs...")
    build_jobs_excel(jobs_sorted)
    build_tailored_excel(jobs_sorted[:50])
    build_tracker_excel(jobs_sorted)
    build_interview_excel(jobs_sorted)
    print("\n🎉 All Excel files updated with premium formatting!")




