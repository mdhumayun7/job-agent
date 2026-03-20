import openpyxl
import json
import os
from openpyxl.styles import PatternFill, Font, Alignment
from datetime import datetime

TRACKER_FILE = r"C:\Users\HP\Desktop\job_agent\output\apply_tracker.xlsx"
JOBS_FILE = r"C:\Users\HP\Desktop\job_agent\output\jobs_found.json"

STATUS_COLORS = {"Not Applied":"E2E8F0","Applied":"BFDBFE","Seen":"FEF08A","Interview":"86EFAC","Rejected":"FCA5A5","Offer":"4ADE80","Follow Up":"FDE68A"}

def create_tracker():
    with open(JOBS_FILE, "r", encoding="utf-8") as jf:
        jobs = json.load(jf)

    if os.path.exists(TRACKER_FILE):
        wb = openpyxl.load_workbook(TRACKER_FILE)
        ws = wb.active
        existing = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[1]: existing.add(str(row[1])[:50])
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Apply Tracker"
        existing = set()
        headers = ["Date Added","Job Title","Company","Location","Salary","Source","Status","Applied Date","Interview Date","Follow Up Date","Notes","Apply Link"]
        hf = PatternFill("solid", fgColor="1E293B")
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = hf
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")

    new_count = 0
    for job in jobs:
        title_key = str(job.get("title",""))[:50]
        if title_key in existing: continue
        row = ws.max_row + 1
        ws.cell(row=row, column=1, value=datetime.now().strftime("%d-%m-%Y"))
        ws.cell(row=row, column=2, value=job.get("title","")[:100])
        ws.cell(row=row, column=3, value=job.get("company",""))
        ws.cell(row=row, column=4, value=job.get("location",""))
        ws.cell(row=row, column=5, value=job.get("salary","Not mentioned"))
        ws.cell(row=row, column=6, value=job.get("source",""))
        sc = ws.cell(row=row, column=7, value="Not Applied")
        sc.fill = PatternFill("solid", fgColor="E2E8F0")
        ws.cell(row=row, column=8, value="")
        ws.cell(row=row, column=9, value="")
        ws.cell(row=row, column=10, value="")
        ws.cell(row=row, column=11, value="")
        ws.cell(row=row, column=12, value=job.get("apply_url",""))
        new_count += 1

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len+4, 50)

    wb.save(TRACKER_FILE)
    print(f"Tracker updated! {new_count} new jobs added.")
    print(f"File: {TRACKER_FILE}")

if __name__ == "__main__":
    create_tracker()
